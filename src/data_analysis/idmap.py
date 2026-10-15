"""Gene ID mapping: mouse symbol <-> Ensembl, mouse -> rat/human 1:1 orthologs."""

from __future__ import annotations

import json
import logging
import re
from pathlib import Path
from typing import Iterable

import openpyxl
import pandas as pd

log = logging.getLogger(__name__)

TAXID = {"mouse": 10090, "rat": 10116, "human": 9606}
ENSEMBL_PREFIX = {"mouse": "ENSMUSG", "rat": "ENSRNOG", "human": "ENSG"}
MHC_PATTERN = re.compile(r"^H2-")  # mouse MHC genes have no 1:1 rat RT1-* / human HLA-* match
# Clustered protocadherins expanded independently per lineage; identical symbols are not 1:1 orthologs.
NO_SYMBOL_FALLBACK = re.compile(r"^Pcdh[abg]\d")

# Older Ensembl releases used different IDs for some genes. Walker et al. 2018 (GSE110344) counts
# predate the Lilrb4 -> Lilrb4a/Lilrb4b split, so Lilrb4a is stored under its former gene ID.
RETIRED_MOUSE_IDS = {"Lilrb4a": ["ENSMUSG00000062593"]}


# --------------------------------------------------------------------------------------------
# mygene with an on-disk cache, so reruns are offline and reproducible
# --------------------------------------------------------------------------------------------

class MyGeneCache:
    """Thin wrapper around mygene.querymany that stores every hit in a JSON file."""

    def __init__(self, cache_file: str | Path):
        self.cache_file = Path(cache_file)
        self._store: dict[str, list[dict]] = {}
        if self.cache_file.is_file():
            self._store = json.loads(self.cache_file.read_text(encoding="utf-8"))
        self._client = None

    def _key(self, query: str, scopes: str, species: str | int, fields: str) -> str:
        return f"{scopes}|{species}|{fields}|{query}"

    def querymany(self, queries: Iterable[str], scopes: str, species: str | int, fields: str) -> dict[str, list[dict]]:
        """Return {query: [hits...]} (empty list when mygene finds nothing)."""
        queries = list(dict.fromkeys(str(q) for q in queries))
        todo = [q for q in queries if self._key(q, scopes, species, fields) not in self._store]
        if todo:
            if self._client is None:
                import mygene
                self._client = mygene.MyGeneInfo()
            result = self._client.querymany(
                todo, scopes=scopes, species=species, fields=fields, returnall=True, verbose=False
            )
            hits: dict[str, list[dict]] = {q: [] for q in todo}
            for hit in result["out"]:
                if not hit.get("notfound"):
                    hits[str(hit["query"])].append({k: v for k, v in hit.items() if k != "query"})
            for q, hs in hits.items():
                self._store[self._key(q, scopes, species, fields)] = hs
            self.cache_file.parent.mkdir(parents=True, exist_ok=True)
            self.cache_file.write_text(json.dumps(self._store, indent=1, sort_keys=True), encoding="utf-8")
        return {q: self._store[self._key(q, scopes, species, fields)] for q in queries}


def _ensembl_ids(hit: dict, species: str) -> list[str]:
    """Extract Ensembl gene IDs of the reference strain for `species` from a mygene hit."""
    ens = hit.get("ensembl") or []
    if isinstance(ens, dict):
        ens = [ens]
    ids = [e["gene"] for e in ens if isinstance(e, dict) and "gene" in e]
    return sorted({i for i in ids if i.startswith(ENSEMBL_PREFIX[species])})


# --------------------------------------------------------------------------------------------
# 1. mouse symbol <-> Ensembl
# --------------------------------------------------------------------------------------------

def read_deglists_id_table(deglists: str | Path) -> pd.DataFrame:
    """Collect unique (ensembl_gene_id, symbol) pairs from every sheet of GSE228031_DEGlists.xlsx."""
    wb = openpyxl.load_workbook(deglists, read_only=True)
    pairs: set[tuple[str, str]] = set()
    for name in wb.sheetnames:
        rows = wb[name].iter_rows(values_only=True)
        header = list(next(rows))
        if "ensembl_gene_id" not in header:
            continue  # SUMMARY sheet
        i_ens, i_sym = header.index("ensembl_gene_id"), header.index("external_gene_name")
        for row in rows:
            if row[i_ens] and row[i_sym]:
                pairs.add((str(row[i_ens]).strip(), str(row[i_sym]).strip()))
    wb.close()
    table = pd.DataFrame(sorted(pairs), columns=["ensembl_gene_id", "symbol"])
    if table["ensembl_gene_id"].duplicated().any():
        raise ValueError("an Ensembl ID maps to more than one symbol in DEGlists")
    return table


def load_mouse_id_table(deglists: str | Path, cache_dir: str | Path) -> pd.DataFrame:
    """Return the DEGlists ID table, cached as parquet (reading all sheets takes ~30 s)."""
    cache = Path(cache_dir) / "mouse_symbol_ensembl.parquet"
    if cache.is_file():
        return pd.read_parquet(cache)
    table = read_deglists_id_table(deglists)
    cache.parent.mkdir(parents=True, exist_ok=True)
    table.to_parquet(cache, index=False)
    return table


def map_mouse_symbols(genes: Iterable[str], id_table: pd.DataFrame, mg: MyGeneCache) -> pd.DataFrame:
    """Map mouse symbols to Ensembl: DEGlists first, then mygene by symbol and alias.

    Columns: gene, ensembl_gene_id, current_symbol, id_source ('deglists'|'mygene'|'unmapped').
    """
    genes = list(genes)
    sym2ens = id_table.groupby("symbol")["ensembl_gene_id"].agg(list).to_dict()
    records = []
    fallback = []
    for gene in genes:
        ids = sym2ens.get(gene, [])
        if len(ids) == 1:
            records.append({"gene": gene, "ensembl_gene_id": ids[0], "current_symbol": gene, "id_source": "deglists"})
        else:
            if len(ids) > 1:
                log.warning("%s has %d Ensembl IDs in DEGlists; resolving with mygene", gene, len(ids))
            fallback.append(gene)

    hits = mg.querymany(fallback, scopes="symbol,alias", species="mouse", fields="symbol,ensembl.gene")
    for gene in fallback:
        candidates = sorted(hits[gene], key=lambda h: (h.get("symbol") != gene, -h.get("_score", 0)))
        rec = {"gene": gene, "ensembl_gene_id": None, "current_symbol": None, "id_source": "unmapped"}
        for hit in candidates:
            ids = _ensembl_ids(hit, "mouse")
            if ids:
                if len(ids) > 1:
                    log.warning("%s: several mouse Ensembl IDs from mygene %s; using %s", gene, ids, ids[0])
                rec.update(ensembl_gene_id=ids[0], current_symbol=hit.get("symbol"), id_source="mygene")
                break
        if rec["id_source"] == "unmapped":
            log.warning("%s could not be mapped to a mouse Ensembl ID", gene)
        records.append(rec)

    out = pd.DataFrame(records).set_index("gene").loc[genes].reset_index()
    out["retired_ensembl_ids"] = out["gene"].map(lambda g: ";".join(RETIRED_MOUSE_IDS.get(g, [])) or None)
    return out


def mouse_id_candidates(row: pd.Series) -> list[str]:
    """Current plus retired Ensembl IDs for one row of `map_mouse_symbols` output."""
    ids = [row["ensembl_gene_id"]] if isinstance(row["ensembl_gene_id"], str) else []
    if isinstance(row.get("retired_ensembl_ids"), str):
        ids += row["retired_ensembl_ids"].split(";")
    return ids


# --------------------------------------------------------------------------------------------
# 2-3. mouse -> rat / human 1:1 orthologs
# --------------------------------------------------------------------------------------------

def map_orthologs(mouse: pd.DataFrame, target: str, mg: MyGeneCache) -> pd.DataFrame:
    """Map mouse genes (columns gene, ensembl_gene_id) to 1:1 orthologs in `target` ('rat'|'human').

    Uses HomoloGene groups from mygene. When HomoloGene has no ortholog, falls back to a target
    gene with the same symbol (rat: as is, human: upper case), status 'symbol_match'.
    Genes are excluded when they are mouse MHC (H2-*), have no ortholog, or have several
    orthologs in the target species.
    Columns: gene, {target}_symbol, {target}_ensembl, {target}_status.
    """
    if target not in ("rat", "human"):
        raise ValueError(f"unsupported target species {target!r}")
    taxid = TAXID[target]

    mapped = mouse.dropna(subset=["ensembl_gene_id"])
    homolog = mg.querymany(mapped["ensembl_gene_id"], scopes="ensembl.gene", species="mouse", fields="homologene")

    target_entrez: dict[str, str] = {}
    status: dict[str, str] = {}
    for gene, ens in zip(mouse["gene"], mouse["ensembl_gene_id"]):
        if MHC_PATTERN.match(gene):
            status[gene] = "excluded_mhc"
            continue
        if not isinstance(ens, str):
            status[gene] = "no_mouse_id"
            continue
        groups = [h.get("homologene", {}) for h in homolog[ens]]
        entrez = sorted({str(g[1]) for grp in groups for g in grp.get("genes", []) if g[0] == taxid})
        if not entrez:
            status[gene] = "no_ortholog"
        elif len(entrez) > 1:
            status[gene] = "one_to_many"
        else:
            status[gene] = "one_to_one"
            target_entrez[gene] = entrez[0]

    # Symbol fallback for genes HomoloGene does not cover.
    current = dict(zip(mouse["gene"], mouse["current_symbol"]))
    fallback = {
        g: (current[g] if target == "rat" else current[g].upper())
        for g, s in status.items()
        if s == "no_ortholog" and isinstance(current.get(g), str) and not NO_SYMBOL_FALLBACK.match(g)
    }
    by_symbol = mg.querymany(fallback.values(), scopes="symbol", species=taxid, fields="symbol,entrezgene")
    for gene, symbol in fallback.items():
        exact = [h for h in by_symbol[symbol] if str(h.get("symbol", "")).lower() == symbol.lower()]
        if len(exact) == 1 and exact[0].get("entrezgene"):
            status[gene] = "symbol_match"
            target_entrez[gene] = str(exact[0]["entrezgene"])

    info = mg.querymany(target_entrez.values(), scopes="entrezgene", species=taxid, fields="symbol,ensembl.gene")
    rows = []
    for gene in mouse["gene"]:
        row = {"gene": gene, f"{target}_symbol": None, f"{target}_ensembl": None, f"{target}_status": status[gene]}
        if gene in target_entrez:
            hits = info[target_entrez[gene]]
            if hits:
                row[f"{target}_symbol"] = hits[0].get("symbol")
                ids = _ensembl_ids(hits[0], target)
                row[f"{target}_ensembl"] = ";".join(ids) if ids else None
            else:
                row[f"{target}_status"] = "no_target_info"
        rows.append(row)
    return pd.DataFrame(rows)


def build_cas_mapping(signature: pd.DataFrame, id_table: pd.DataFrame, mg: MyGeneCache) -> pd.DataFrame:
    """Combine mouse Ensembl IDs and rat/human orthologs for every CAS gene."""
    mouse = map_mouse_symbols(signature["gene"], id_table, mg)
    out = signature[["gene", "direction"]].merge(mouse, on="gene", how="left")
    for target in ("rat", "human"):
        out = out.merge(map_orthologs(mouse, target, mg), on="gene", how="left")
    return out
